from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def set_cell_border(cell, **kwargs):
    """设置单元格边框"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = 'w:{}'.format(edge)
            element = tcPr.find(qn(tag))
            if element is None:
                element = docx.oxml.OxmlElement(tag)
                tcPr.append(element)
            element.set(qn('w:val'), 'single')
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))

def add_heading(doc, text, size=14, bold=True, color='#1a73e8', space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16))
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(6)
    return p

def add_bullet(doc, text, indent=0.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    p.paragraph_format.space_after = Pt(2)
    return p

def add_text(doc, text, size=9.5, bold=False, color='#333333'):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16))
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    p.paragraph_format.space_after = Pt(2)
    return p

def add_two_column(doc, left_text, right_text, size=9.5, bold=False):
    p = doc.add_paragraph()
    run1 = p.add_run(left_text)
    run1.font.size = Pt(size)
    run1.font.bold = bold
    run1.font.name = 'Microsoft YaHei'
    run1._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    p.add_run(' ' * 80).font.size = Pt(size)
    run2 = p.add_run(right_text)
    run2.font.size = Pt(size)
    run2.font.name = 'Microsoft YaHei'
    run2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    p.paragraph_format.space_after = Pt(2)
    return p

# 创建文档
doc = Document()

# 设置全局样式
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
style.font.size = Pt(9.5)
style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# 页面边距
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

# ====== 标题 ======
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run('陈铖')
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x1a, 0x73, 0xe8)
title_run.font.name = 'Microsoft YaHei'
title_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
title.paragraph_format.space_after = Pt(2)

# 求职意向
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = subtitle.add_run('AI CSM实习生 | AI解决方案实习生')
sub_run.font.size = Pt(12)
sub_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
sub_run.font.name = 'Microsoft YaHei'
sub_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
subtitle.paragraph_format.space_after = Pt(4)

# 联系方式
contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run = contact.add_run('手机：150-0000-0000 | 邮箱：chencheng@example.com | 微信：chencheng_ai')
contact_run.font.size = Pt(9)
contact_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
contact_run.font.name = 'Microsoft YaHei'
contact_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
contact.paragraph_format.space_after = Pt(8)

# ====== 教育背景 ======
add_heading(doc, '教育背景', size=13)
add_two_column(doc, '中国民航大学 | 计算机科学与技术 | 本科', '2022.09 - 2026.06', bold=True)
add_bullet(doc, '主修课程：数据结构、数据库原理、软件工程、计算机网络、人工智能导论')

# ====== 实习经历 ======
add_heading(doc, '实习经历', size=13)
add_two_column(doc, '拓尔思信息技术股份有限公司 | AI应用开发实习生', '2024.07 - 至今', bold=True)
add_bullet(doc, '参与AI Agent应用开发项目，协助团队完成需求调研、方案设计和Demo搭建')
add_bullet(doc, '使用Python+FastAPI开发AI服务后端，对接LLM API实现智能问答、内容生成等功能')
add_bullet(doc, '编写项目文档和技术方案，包括需求分析文档、接口文档和部署说明')
add_bullet(doc, '协助进行用户反馈收集和效果复盘，整理问题清单并推动优化迭代')

# ====== 项目经历 ======
add_heading(doc, '项目经历', size=13)

add_two_column(doc, '轻量化AI营销系统', '个人项目', bold=True)
add_bullet(doc, '基于AI大模型能力，设计面向中小企业的轻量化营销解决方案，覆盖文案生成、客户画像、效果追踪三大场景')
add_bullet(doc, '使用FastAPI搭建后端服务，接入LLM API实现多场景文案自动生成和个性化内容推荐')
add_bullet(doc, '用Vue开发前端界面，实现营销仪表盘、文案管理和客户管理模块')
add_bullet(doc, '独立完成从需求分析到方案设计、开发实现到效果验证的全流程，积累AI产品落地经验')

add_two_column(doc, 'AI面试官&简历评估系统', '课程/项目', bold=True)
add_bullet(doc, '针对企业招聘场景，设计AI面试官产品方案，定义产品功能和用户体验流程')
add_bullet(doc, '基于RAG架构构建面试知识库，使用Chroma向量数据库存储和检索面试题库')
add_bullet(doc, '设计Prompt策略实现多轮对话、智能追问和评分反馈，支持不同岗位的面试模板')
add_bullet(doc, '输出PRD文档和原型设计，用Axure绘制产品交互流程图')

add_two_column(doc, 'BOSS直聘智能投递工具', '个人工具开发', bold=True)
add_bullet(doc, '为解决批量投递效率低的问题，设计自动化投递工具，实现职位筛选、自动沟通和数据追踪')
add_bullet(doc, '使用DrissionPage进行网页自动化操作，实现职位抓取、过滤和自动点击沟通')
add_bullet(doc, '设计规则飞轮系统：通过人工标注→规则挖掘→自动更新的闭环，持续提升筛选准确率')
add_bullet(doc, '开发Tkinter GUI界面，支持配置管理、实时状态监控和投递数据统计')

# ====== 技能与工具 ======
add_heading(doc, '技能与工具', size=13)
add_bullet(doc, '技术能力：会用Python进行数据处理和自动化脚本开发，了解FastAPI后端开发，会用SQL进行数据查询，了解API调用和LLM接入')
add_bullet(doc, 'AI工具：实际使用过ChatGPT、Claude、Cursor、Coze、Dify、豆包等工具，有Prompt Engineering调优经验')
add_bullet(doc, '产品工具：会用Axure画原型、Figma做UI设计、Notion写文档、飞书进行协作')
add_bullet(doc, '文档能力：能独立输出需求分析文档、解决方案文档、项目复盘材料和演示PPT')

# ====== 个人优势 ======
add_heading(doc, '个人优势', size=13)
add_bullet(doc, '对AI落地有实际项目经验，理解从需求分析到方案交付的完整流程')
add_bullet(doc, '具备技术基础，能与研发团队顺畅沟通，快速理解技术方案和限制')
add_bullet(doc, '动手能力强，习惯用Demo和原型验证想法，愿意尝试新工具和新方法')
add_bullet(doc, '沟通协作意识好，在实习期间与产品、技术、客户成功等多角色配合推进项目')
add_bullet(doc, '实习时间稳定，可保证每周5天、连续实习1年')

# 保存
doc.save(r'D:\陈铖-AI CSM实习生-飞书.docx')
print('简历已生成: D:\陈铖-AI CSM实习生-飞书.docx')
